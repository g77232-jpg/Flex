#!/usr/bin/env python3
# Copyright (c) Gbickham elventh110ur production.
from __future__ import annotations

import argparse
import datetime as dt
import importlib.util
import os
from pathlib import Path
from typing import Iterable, Sequence


DEFAULT_KEYWORDS = [
    "11th hour",
    "11th-hour",
    "11th_hour",
    "11th hour production",
    "cohort pilot",
    "cohert pilot",
    "instructor workbook",
    "business plan",
    "code of conduct",
    "codes of conduct",
]
DEFAULT_EXTENSIONS = [".docx", ".pdf", ".txt", ".md", ".rtf"]
DEFAULT_ROOTS = ["."]
EXTRA_ROOTS = ["/media", "/mnt", "/run/media"]
COPYRIGHT_NOTICE = "Copyright (c) Gbickham elventh110ur production."
IGNORE_DIRS = {
    ".git",
    ".hg",
    ".svn",
    "__pycache__",
    "node_modules",
    ".venv",
    "venv",
}


def module_available(module_name: str) -> bool:
    return importlib.util.find_spec(module_name) is not None


def normalize_roots(roots: Sequence[str]) -> list[Path]:
    expanded = []
    for root in roots:
        candidate = Path(root).expanduser()
        if candidate.exists():
            expanded.append(candidate.resolve())
    return expanded


def iter_candidate_files(
    roots: Iterable[Path],
    extensions: set[str],
    max_size_bytes: int,
) -> Iterable[Path]:
    for root in roots:
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames if d not in IGNORE_DIRS]
            for filename in filenames:
                path = Path(dirpath, filename)
                if path.suffix.lower() not in extensions:
                    continue
                try:
                    size = path.stat().st_size
                except OSError:
                    continue
                if size > max_size_bytes:
                    continue
                yield path


def matches_keywords(text: str, keywords: Sequence[str]) -> bool:
    lowered = text.lower()
    return any(keyword.lower() in lowered for keyword in keywords)


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx_file(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)


def read_pdf_file(path: Path) -> str:
    import PyPDF2

    text_chunks = []
    with path.open("rb") as handle:
        reader = PyPDF2.PdfReader(handle)
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def gather_documents(
    roots: Sequence[Path],
    keywords: Sequence[str],
    extensions: Sequence[str],
    max_size_mb: int,
) -> tuple[list[tuple[Path, str]], list[str]]:
    extensions_set = {ext.lower() for ext in extensions}
    max_size_bytes = max_size_mb * 1024 * 1024
    found: list[tuple[Path, str]] = []
    skipped: list[str] = []

    for path in iter_candidate_files(roots, extensions_set, max_size_bytes):
        if not matches_keywords(path.name, keywords):
            if path.suffix.lower() in {".txt", ".md", ".rtf"}:
                content = read_text_file(path)
                if not matches_keywords(content, keywords):
                    continue
            else:
                continue

        try:
            if path.suffix.lower() == ".docx":
                if not module_available("docx"):
                    skipped.append(f"Missing python-docx for {path}")
                    continue
                content = read_docx_file(path)
            elif path.suffix.lower() == ".pdf":
                if not module_available("PyPDF2"):
                    skipped.append(f"Missing PyPDF2 for {path}")
                    continue
                content = read_pdf_file(path)
            else:
                content = read_text_file(path)
        except Exception as exc:  # noqa: BLE001
            skipped.append(f"Failed to read {path}: {exc}")
            continue

        if content.strip():
            found.append((path, content))

    return found, skipped


def build_docx(
    output_path: Path,
    sources: Sequence[tuple[Path, str]],
    roots: Sequence[Path],
    keywords: Sequence[str],
    skipped: Sequence[str],
) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Business Plan Compilation", 0)
    document.add_paragraph(COPYRIGHT_NOTICE)
    document.add_paragraph(f"Generated: {dt.datetime.now(dt.timezone.utc).isoformat()}")
    document.add_paragraph(f"Roots scanned: {', '.join(str(root) for root in roots)}")
    document.add_paragraph(f"Keywords: {', '.join(keywords)}")

    if not sources:
        document.add_paragraph("No matching documents were found.")
    else:
        for path, content in sources:
            document.add_heading(path.name, level=1)
            document.add_paragraph(f"Source: {path}")
            for line in content.splitlines():
                document.add_paragraph(line)

    if skipped:
        document.add_heading("Skipped Items", level=1)
        for item in skipped:
            document.add_paragraph(item)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Scan for business-plan-related documents (including SD card paths) "
            "and compile them into a single DOCX."
        )
    )
    parser.add_argument(
        "--roots",
        nargs="*",
        default=DEFAULT_ROOTS + EXTRA_ROOTS,
        help="Root directories to scan. Defaults to current directory plus common SD card mounts.",
    )
    parser.add_argument(
        "--keywords",
        nargs="*",
        default=DEFAULT_KEYWORDS,
        help="Keywords to match in filenames or document text.",
    )
    parser.add_argument(
        "--extensions",
        nargs="*",
        default=DEFAULT_EXTENSIONS,
        help="File extensions to include.",
    )
    parser.add_argument(
        "--max-size-mb",
        type=int,
        default=25,
        help="Maximum file size (MB) per document.",
    )
    parser.add_argument(
        "--output",
        default="output/business_plan_compilation.docx",
        help="Output DOCX path.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    roots = normalize_roots(args.roots)

    if not module_available("docx"):
        raise SystemExit(
            "python-docx is required to generate the DOCX. "
            "Install it with: pip install python-docx"
        )

    sources, skipped = gather_documents(
        roots=roots,
        keywords=args.keywords,
        extensions=args.extensions,
        max_size_mb=args.max_size_mb,
    )

    output_path = Path(args.output).expanduser()
    build_docx(output_path, sources, roots, args.keywords, skipped)
    print(f"Created {output_path} with {len(sources)} document(s).")
    if skipped:
        print("Skipped:")
        for item in skipped:
            print(f"- {item}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
